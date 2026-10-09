"""Evaluate the pipeline on 50 questions and write results to Excel.

Output workbook has two sheets:
  - "Per Question" - one row per question with all details
  - "Summary" - aggregate metrics for reporting

Requires openpyxl:  pip install openpyxl
"""

import sys, os
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "scripts", "retrieval"))
sys.path.insert(0, os.path.join(HERE, "..", "chatbot", "scripts"))

from scripts.retrieval.search        import HybridRetriever
from scripts.retrieval.translator    import Translator
from scripts.retrieval.two_stage     import two_stage_search
from scripts.retrieval.extractor_map import EXTRACTORS_FOR_INTENT
from chatbot.scripts.intent        import detect_intents

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ---------------------------------------------------------------
# Ground truth - paste your final 50 here.
# Start with the 15 already-verified, add 35 new below.
# ---------------------------------------------------------------

GROUND_TRUTH = [
    # === EXISTING 15 (already verified) ===
    {
        "id":                "Q1",
        "question":          "বয়স্ক ভাতা পেতে সর্বনিম্ন বয়স কত হতে হবে?",
        "ministry":          "Social Welfare",
        "acceptable_ids":    [
            "SocialMin_OldAgeAll_2013_00_Policy_v1.pdf_C0006",
            "SocialMin_OldAgeAll_2013_00_Policy_v1.pdf_C0009",
        ],
        "expected_value":    "65",
        "expected_extractor":"extract_age",
        "category":          "basic_extraction",
        "translate":         False,
    },
    # ... rest of your 15 curated questions here
    # [paste your existing 15 from evaluate_metrics.py]

    # === NEW 35 (fill in from your hand-verified list) ===
    # {
    #     "id":                "Q16",
    #     "question":          "...",
    #     ...
    # },
]


# ---------------------------------------------------------------
# Metric computation
# ---------------------------------------------------------------

def recall_at_k(retriever, question, ministry, acceptable_ids, k):
    if not acceptable_ids:
        return None
    results = retriever.search(question, top_k=k, ministry=ministry)
    return bool({r["id"] for r in results} & set(acceptable_ids))


def rank_of_first_hit(results, acceptable_ids):
    for i, r in enumerate(results, 1):
        if r["id"] in acceptable_ids:
            return i
    return None


def check_extracted_value(extracted, expected_extractor, expected_value):
    if expected_extractor is None or expected_value is None:
        return None
    if expected_extractor not in extracted:
        return False
    actual = extracted[expected_extractor].get("value")
    actual_str = ",".join(str(x) for x in actual) if isinstance(actual, list) else str(actual)
    actual_norm = actual_str.replace(",", "").replace(" ", "")
    if isinstance(expected_value, list):
        return any(e.replace(",", "").replace(" ", "") in actual_norm
                   for e in expected_value)
    return str(expected_value).replace(",", "").replace(" ", "") in actual_norm


# ---------------------------------------------------------------
# Main
# ---------------------------------------------------------------

def main():
    print("[startup] loading retriever ...")
    r  = HybridRetriever()
    tr = Translator()
    print("[startup] ready\n")

    rows = []
    for tc in GROUND_TRUTH:
        question   = tc["question"]
        translated = tr.to_bengali(question) if tc.get("translate") else question
        intents    = detect_intents(translated)

        # Recall metrics (before re-ranking)
        r10  = recall_at_k(r, translated, tc["ministry"], tc["acceptable_ids"], 10)
        r50  = recall_at_k(r, translated, tc["ministry"], tc["acceptable_ids"], 50)
        r100 = recall_at_k(r, translated, tc["ministry"], tc["acceptable_ids"], 100)

        # Full pipeline
        results = two_stage_search(
            retriever             = r,
            question              = translated,
            intents               = intents,
            extractors_for_intent = EXTRACTORS_FOR_INTENT,
            ministry              = tc["ministry"],
            return_k              = 10,
        )

        if tc.get("category") == "low_conf_deferral":
            top       = results[0] if results else None
            hit1      = top and top.get("confidence_tier") == "low"
            hit3, hit5 = hit1, hit1
            rank      = 1 if hit1 else None
            extraction_ok = None
            top_id    = top["id"] if top else ""
        elif tc.get("category") == "out_of_scope":
            top       = results[0] if results else None
            hit1      = top and top.get("confidence_tier") == "out_of_scope"
            hit3, hit5 = hit1, hit1
            rank      = 1 if hit1 else None
            extraction_ok = None
            top_id    = "OOS" if hit1 else (top["id"] if top else "")
        else:
            rank      = rank_of_first_hit(results, tc["acceptable_ids"])
            hit1      = rank == 1
            hit3      = rank is not None and rank <= 3
            hit5      = rank is not None and rank <= 5
            top_id    = results[0]["id"] if results else ""
            if hit1:
                top_extracted = results[0].get("extracted", {})
                extraction_ok = check_extracted_value(
                    top_extracted,
                    tc.get("expected_extractor"),
                    tc.get("expected_value"),
                )
            else:
                extraction_ok = None

        rr = 1.0 / rank if rank else 0.0

        rows.append({
            "id":            tc["id"],
            "category":      tc.get("category", ""),
            "ministry":      tc["ministry"],
            "question":      question,
            "translated":    translated if translated != question else "",
            "intents":       ", ".join(intents) if intents else "",
            "expected_ids":  ", ".join(tc["acceptable_ids"]) if tc["acceptable_ids"] else "",
            "top_id":        top_id,
            "rank":          rank,
            "hit1":          hit1,
            "hit3":          hit3,
            "hit5":          hit5,
            "reciprocal":    round(rr, 4),
            "r10":           r10,
            "r50":           r50,
            "r100":          r100,
            "extraction_ok": extraction_ok,
            "expected_val":  tc.get("expected_value", ""),
        })
        print(f"  {tc['id']:<8} rank={str(rank):>4} {question[:60]}")

    # -----------------------------------------------------------
    # Write to Excel
    # -----------------------------------------------------------
    wb = Workbook()

    ws = wb.active
    ws.title = "Per Question"

    header = [
        "ID", "Category", "Ministry", "Question", "Translated",
        "Intents", "Expected IDs", "Top-1 ID", "Rank",
        "Hit@1", "Hit@3", "Hit@5", "Reciprocal",
        "R@10", "R@50", "R@100", "Extraction OK", "Expected Value",
    ]
    ws.append(header)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")

    for row in rows:
        ws.append([
            row["id"], row["category"], row["ministry"],
            row["question"], row["translated"],
            row["intents"], row["expected_ids"], row["top_id"],
            row["rank"] if row["rank"] is not None else "",
            row["hit1"], row["hit3"], row["hit5"], row["reciprocal"],
            row["r10"], row["r50"], row["r100"],
            row["extraction_ok"] if row["extraction_ok"] is not None else "-",
            row["expected_val"] if not isinstance(row["expected_val"], list)
                else ", ".join(row["expected_val"]),
        ])

    for col_letter, width in zip("ABCDEFGHIJKLMNOPQR",
                                  [8, 20, 20, 60, 40, 20, 40, 40, 6,
                                   8, 8, 8, 12, 8, 8, 8, 14, 20]):
        ws.column_dimensions[col_letter].width = width

    # ---- Summary sheet ----
    ws2 = wb.create_sheet("Summary")
    n = len(rows)

    with_ids = [r for r in rows if r["r10"] is not None]
    n_ids    = len(with_ids)

    summary = [
        ["Metric", "Value", "Details"],
        ["Total questions", n, ""],
        ["", "", ""],
        ["Retrieval - Recall@10",
         round(sum(1 for r in with_ids if r["r10"]) / n_ids, 3) if n_ids else "n/a",
         f"{n_ids} questions with acceptable_ids"],
        ["Retrieval - Recall@50",
         round(sum(1 for r in with_ids if r["r50"]) / n_ids, 3) if n_ids else "n/a", ""],
        ["Retrieval - Recall@100",
         round(sum(1 for r in with_ids if r["r100"]) / n_ids, 3) if n_ids else "n/a", ""],
        ["", "", ""],
        ["Full pipeline - MRR",
         round(sum(r["reciprocal"] for r in rows) / n, 3), ""],
        ["Full pipeline - Hit@1",
         round(sum(1 for r in rows if r["hit1"]) / n, 3),
         f"{sum(1 for r in rows if r['hit1'])} / {n}"],
        ["Full pipeline - Hit@3",
         round(sum(1 for r in rows if r["hit3"]) / n, 3),
         f"{sum(1 for r in rows if r['hit3'])} / {n}"],
        ["Full pipeline - Hit@5",
         round(sum(1 for r in rows if r["hit5"]) / n, 3),
         f"{sum(1 for r in rows if r['hit5'])} / {n}"],
        ["", "", ""],
    ]

    ext_checked = [r for r in rows if r["extraction_ok"] is not None]
    if ext_checked:
        ext_ok = sum(1 for r in ext_checked if r["extraction_ok"])
        summary.append([
            "Extraction Accuracy",
            round(ext_ok / len(ext_checked), 3),
            f"{ext_ok} / {len(ext_checked)} on top-1 correct retrievals",
        ])

    summary.append(["", "", ""])
    summary.append(["Evaluated at", datetime.now().strftime("%Y-%m-%d %H:%M"), ""])

    for row in summary:
        ws2.append(row)

    for cell in ws2[1]:
        cell.font = Font(bold=True)

    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 15
    ws2.column_dimensions["C"].width = 50

    # Save
    outpath = os.path.join(HERE, f"evaluation_results_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")
    wb.save(outpath)
    print(f"\nExcel report saved: {outpath}")


if __name__ == "__main__":
    main()